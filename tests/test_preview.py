"""F6-2: 在线预览 测试"""

import io
import json
import os
import tempfile

import pytest
import bcrypt

from models import db as _db, File, Folder, Project, User


@pytest.fixture
def app():
    db_fd, db_path = tempfile.mkstemp(suffix='.db')
    db_uri = 'sqlite:///' + db_path

    key_fd, key_path = tempfile.mkstemp(suffix='.setup_key')
    with open(key_path, 'w', encoding='utf-8') as kf:
        kf.write('test-key')

    from app import create_app
    app = create_app(config_overrides={
        'SQLALCHEMY_DATABASE_URI': db_uri,
        'TESTING': True,
        'SETUP_KEY_FILE': key_path,
    })

    yield app

    with app.app_context():
        _db.session.remove()
        _db.engine.dispose()
    os.close(db_fd)
    os.close(key_fd)
    try:
        os.unlink(db_path)
    except PermissionError:
        pass
    try:
        os.unlink(key_path)
    except PermissionError:
        pass


@pytest.fixture
def client(app):
    return app.test_client()


def _login_as(app, username, password, role):
    with app.app_context():
        pw_hash = bcrypt.hashpw(password.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')
        user = User(
            username=username,
            display_name=username.capitalize(),
            password_hash=pw_hash,
            role=role,
            is_active=True,
            created_ip='127.0.0.1',
        )
        _db.session.add(user)
        _db.session.commit()

    client = app.test_client()
    client.get('/login')
    with client.session_transaction() as sess:
        token = sess.get('csrf_token')
    client.post('/login', data={
        'csrf_token': token,
        'username': username,
        'password': password,
    })
    return client


def _get_csrf(client):
    with client.session_transaction() as sess:
        return sess.get('csrf_token', '')


def _upload_file(client, content, filename, proj_id, folder_id, **extra):
    data = {
        'file': (io.BytesIO(content), filename),
        'project_id': str(proj_id),
        'folder_id': str(folder_id),
        'csrf_token': _get_csrf(client),
    }
    data.update(extra)
    return client.post('/api/files/upload', data=data)


def _seed_project_and_folder(app, project_model='P6T1', project_name='预览测试'):
    with app.app_context():
        admin = User.query.filter_by(username='admin').first()
        proj = Project(model=project_model, name=project_name)
        _db.session.add(proj)
        _db.session.flush()

        root = Folder(
            name=f'{project_model}_{project_name}',
            is_project_root=True,
            project_id=proj.id,
            parent_id=None,
            created_by=admin.id if admin else 1,
        )
        _db.session.add(root)
        _db.session.commit()
        return proj.id, root.id


def _get_file_id(app, filename):
    with app.app_context():
        f = File.query.filter_by(original_filename=filename, is_current=True).first()
        return f.id if f else None


class TestPreviewAPI:

    def test_preview_nonexistent_file(self, app):
        """预览不存在的文件返回 400"""
        client = _login_as(app, 'admin', 'Admin@Pass1', 'admin')
        resp = client.get('/api/files/99999/preview')
        assert resp.status_code == 400
        data = json.loads(resp.data)
        assert data['success'] is False

    def test_preview_text_file(self, app):
        """TXT 文件返回文本内容"""
        client = _login_as(app, 'admin', 'Admin@Pass1', 'admin')
        proj_id, folder_id = _seed_project_and_folder(app)

        _upload_file(client, 'Hello\nWorld\n测试'.encode('utf-8'), 'note.txt', proj_id, folder_id)
        file_id = _get_file_id(app, 'note.txt')

        resp = client.get(f'/api/files/{file_id}/preview')
        assert resp.status_code == 200
        data = json.loads(resp.data)
        assert data['success'] is True
        assert data['type'] == 'text'
        assert 'Hello' in data['content']

    def test_preview_csv_file(self, app):
        """CSV 文件返回 headers 和 rows"""
        client = _login_as(app, 'admin', 'Admin@Pass1', 'admin')
        proj_id, folder_id = _seed_project_and_folder(app)

        csv_content = 'Name,Age,City\nAlice,30,Beijing\nBob,25,Shanghai'
        _upload_file(client, csv_content.encode('utf-8'), 'data.csv', proj_id, folder_id)
        file_id = _get_file_id(app, 'data.csv')

        resp = client.get(f'/api/files/{file_id}/preview')
        assert resp.status_code == 200
        data = json.loads(resp.data)
        assert data['success'] is True
        assert data['type'] == 'csv'
        assert data['headers'] == ['Name', 'Age', 'City']
        assert len(data['rows']) == 2

    def test_preview_image_returns_stream_url(self, app):
        """图片文件返回 stream 类型和 stream_url"""
        client = _login_as(app, 'admin', 'Admin@Pass1', 'admin')
        proj_id, folder_id = _seed_project_and_folder(app)

        png_data = (
            b'\x89PNG\r\n\x1a\n' + b'\x00' * 4 +
            b'IHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde'
            + b'\x00' * 4 + b'IDATx\x9cc\xf8\x0f\x00\x00\x01\x01\x00\x05\x18\xd8N'
            + b'\x00' * 4 + b'IEND\xaeB`\x82'
        )
        _upload_file(client, png_data, 'test.png', proj_id, folder_id)
        file_id = _get_file_id(app, 'test.png')

        resp = client.get(f'/api/files/{file_id}/preview')
        assert resp.status_code == 200
        data = json.loads(resp.data)
        assert data['success'] is True
        assert data['type'] == 'stream'
        assert f'/api/files/{file_id}/stream' in data['stream_url']

    def test_stream_image_returns_file(self, app):
        """stream 端点返回图片二进制数据"""
        client = _login_as(app, 'admin', 'Admin@Pass1', 'admin')
        proj_id, folder_id = _seed_project_and_folder(app)

        _upload_file(client, b'fake-png-data', 'photo.jpg', proj_id, folder_id)
        file_id = _get_file_id(app, 'photo.jpg')

        resp = client.get(f'/api/files/{file_id}/stream')
        assert resp.status_code == 200
        assert resp.content_type == 'image/jpeg'
        assert resp.data == b'fake-png-data'

    def test_preview_requires_login(self, app):
        """未登录用户无法访问预览 API"""
        client = app.test_client()
        resp = client.get('/api/files/1/preview')
        assert resp.status_code == 401

    def test_preview_docx_word_file(self, app):
        """Word (.docx) 文件返回 HTML 内容"""
        client = _login_as(app, 'admin', 'Admin@Pass1', 'admin')
        proj_id, folder_id = _seed_project_and_folder(app)

        from docx import Document
        import tempfile as tf
        doc = Document()
        doc.add_heading('测试标题', level=1)
        doc.add_paragraph('第一段内容')
        doc.add_paragraph('第二段内容')
        tmp = tf.NamedTemporaryFile(suffix='.docx', delete=False)
        doc.save(tmp.name)
        tmp.close()

        with open(tmp.name, 'rb') as fh:
            content = fh.read()
        os.unlink(tmp.name)

        _upload_file(client, content, 'report.docx', proj_id, folder_id)
        file_id = _get_file_id(app, 'report.docx')

        resp = client.get(f'/api/files/{file_id}/preview')
        assert resp.status_code == 200
        data = json.loads(resp.data)
        assert data['success'] is True
        assert data['type'] == 'html'
        assert '测试标题' in data['content']
        assert '第一段内容' in data['content']

    def test_preview_xlsx_excel_file(self, app):
        """Excel (.xlsx) 文件返回 HTML 内容"""
        client = _login_as(app, 'admin', 'Admin@Pass1', 'admin')
        proj_id, folder_id = _seed_project_and_folder(app)

        import openpyxl
        import tempfile as tf
        wb = openpyxl.Workbook()
        sh = wb.active
        sh.title = '数据'
        sh.append(['Name', 'Score'])
        sh.append(['Alice', '95'])
        sh.append(['Bob', '87'])
        tmp = tf.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(tmp.name)
        wb.close()
        tmp.close()

        with open(tmp.name, 'rb') as fh:
            content = fh.read()
        os.unlink(tmp.name)

        _upload_file(client, content, 'scores.xlsx', proj_id, folder_id)
        file_id = _get_file_id(app, 'scores.xlsx')

        resp = client.get(f'/api/files/{file_id}/preview')
        assert resp.status_code == 200
        data = json.loads(resp.data)
        assert data['success'] is True
        assert data['type'] == 'html'
        assert 'Alice' in data['content']
        assert 'Score' in data['content']

    def test_preview_code_file(self, app):
        """代码文件返回文本预览"""
        client = _login_as(app, 'admin', 'Admin@Pass1', 'admin')
        proj_id, folder_id = _seed_project_and_folder(app)

        _upload_file(client, b'print("hello")', 'script.py', proj_id, folder_id)
        file_id = _get_file_id(app, 'script.py')

        resp = client.get(f'/api/files/{file_id}/preview')
        assert resp.status_code == 200
        data = json.loads(resp.data)
        assert data['success'] is True
        assert data['type'] == 'text'
        assert 'print("hello")' in data['content']
        assert data.get('filename') == 'script.py'

    def test_preview_markdown_file(self, app):
        """Markdown 文件返回文本预览"""
        client = _login_as(app, 'admin', 'Admin@Pass1', 'admin')
        proj_id, folder_id = _seed_project_and_folder(app)

        _upload_file(client, '# 标题\n内容'.encode('utf-8'), 'readme.md', proj_id, folder_id)
        file_id = _get_file_id(app, 'readme.md')

        resp = client.get(f'/api/files/{file_id}/preview')
        assert resp.status_code == 200
        data = json.loads(resp.data)
        assert data['success'] is True
        assert data['type'] == 'text'
        assert '# 标题' in data['content']

    def test_preview_filename_in_response(self, app):
        """预览响应包含 filename 字段"""
        client = _login_as(app, 'admin', 'Admin@Pass1', 'admin')
        proj_id, folder_id = _seed_project_and_folder(app)

        _upload_file(client, b'data', 'test.txt', proj_id, folder_id)
        file_id = _get_file_id(app, 'test.txt')

        resp = client.get(f'/api/files/{file_id}/preview')
        assert resp.status_code == 200
        data = json.loads(resp.data)
        assert data.get('filename') == 'test.txt'


class TestPreviewPage:
    """预览页面路由测试（新标签页模式）"""

    def test_preview_page_text(self, app):
        """TXT 文件预览页面返回 200 并包含文本内容"""
        client = _login_as(app, 'admin', 'Admin@Pass1', 'admin')
        proj_id, folder_id = _seed_project_and_folder(app)

        _upload_file(client, 'Hello World'.encode('utf-8'), 'readme.txt', proj_id, folder_id)
        file_id = _get_file_id(app, 'readme.txt')

        resp = client.get(f'/preview/{file_id}')
        assert resp.status_code == 200
        assert 'Hello World' in resp.data.decode('utf-8')
        assert '纯文本' in resp.data.decode('utf-8')

    def test_preview_page_image(self, app):
        """图片文件预览页面返回 iframe 加载 stream"""
        client = _login_as(app, 'admin', 'Admin@Pass1', 'admin')
        proj_id, folder_id = _seed_project_and_folder(app)

        png_data = (
            b'\x89PNG\r\n\x1a\n' + b'\x00' * 4 +
            b'IHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde'
            + b'\x00' * 4 + b'IDATx\x9cc\xf8\x0f\x00\x00\x01\x01\x00\x05\x18\xd8N'
            + b'\x00' * 4 + b'IEND\xaeB`\x82'
        )
        _upload_file(client, png_data, 'img.png', proj_id, folder_id)
        file_id = _get_file_id(app, 'img.png')

        resp = client.get(f'/preview/{file_id}')
        assert resp.status_code == 200
        html = resp.data.decode('utf-8')
        assert 'preview-stream' in html
        assert f'/api/files/{file_id}/stream' in html

    def test_preview_page_nonexistent_returns_404(self, app):
        """不存在的文件预览页面返回 404"""
        client = _login_as(app, 'admin', 'Admin@Pass1', 'admin')
        resp = client.get('/preview/99999')
        assert resp.status_code == 404

    def test_preview_page_requires_login(self, app):
        """未登录用户访问预览页面重定向到登录页"""
        client = app.test_client()
        resp = client.get('/preview/1', follow_redirects=False)
        assert resp.status_code in (302, 401)

    def test_preview_page_docx(self, app):
        """Word (.docx) 文件预览页面返回 HTML"""
        client = _login_as(app, 'admin', 'Admin@Pass1', 'admin')
        proj_id, folder_id = _seed_project_and_folder(app)

        from docx import Document
        import tempfile as tf
        doc = Document()
        doc.add_heading('测试', level=1)
        doc.add_paragraph('内容')
        tmp = tf.NamedTemporaryFile(suffix='.docx', delete=False)
        doc.save(tmp.name)
        tmp.close()

        with open(tmp.name, 'rb') as fh:
            content = fh.read()
        os.unlink(tmp.name)

        _upload_file(client, content, 'doc.docx', proj_id, folder_id)
        file_id = _get_file_id(app, 'doc.docx')

        resp = client.get(f'/preview/{file_id}')
        assert resp.status_code == 200
        html = resp.data.decode('utf-8')
        assert '测试' in html
        assert '内容' in html
        assert 'Office 文档' in html

    def test_preview_page_shows_filename(self, app):
        """预览页面显示文件名"""
        client = _login_as(app, 'admin', 'Admin@Pass1', 'admin')
        proj_id, folder_id = _seed_project_and_folder(app)

        _upload_file(client, 'hello'.encode('utf-8'), '我的文件.txt', proj_id, folder_id)
        file_id = _get_file_id(app, '我的文件.txt')

        resp = client.get(f'/preview/{file_id}')
        assert resp.status_code == 200
        html = resp.data.decode('utf-8')
        assert '我的文件.txt' in html

    def test_preview_page_stream_has_fallback(self, app):
        """图片预览页面包含 stream fallback 提示"""
        client = _login_as(app, 'admin', 'Admin@Pass1', 'admin')
        proj_id, folder_id = _seed_project_and_folder(app)

        png_data = (
            b'\x89PNG\r\n\x1a\n' + b'\x00' * 4 +
            b'IHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde'
            + b'\x00' * 4 + b'IDATx\x9cc\xf8\x0f\x00\x00\x01\x01\x00\x05\x18\xd8N'
            + b'\x00' * 4 + b'IEND\xaeB`\x82'
        )
        _upload_file(client, png_data, 'img.png', proj_id, folder_id)
        file_id = _get_file_id(app, 'img.png')

        resp = client.get(f'/preview/{file_id}')
        assert resp.status_code == 200
        html = resp.data.decode('utf-8')
        assert 'stream-fallback' in html
